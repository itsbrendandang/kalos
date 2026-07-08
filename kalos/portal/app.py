"""Kalos Engine portal — a small web view that runs the real engine.

A FastAPI app that, on request, runs the BoTorch optimization (single- and
multi-objective) on a synthetic bioprocess surface and returns the results, plus
a single page that charts them. It is a *viewer over the live engine*, not a
mock: every number comes from an actual BoTorch fit + acquisition.

Run:  python -m kalos.portal   (then open http://127.0.0.1:8050)
Needs the portal extra:  pip install -e ".[portal]"
"""
from __future__ import annotations

import io
import json
import logging
import os
import re
import threading
import time
from pathlib import Path

import numpy as np
import pandas as pd
import torch
from fastapi import FastAPI, File, Form, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import HTMLResponse, JSONResponse
from scipy.stats import spearmanr
from sklearn.decomposition import PCA
from sklearn.preprocessing import StandardScaler
from starlette.concurrency import run_in_threadpool

from kalos import __version__ as ENGINE_VERSION
from kalos.core.conformal import q_from_residuals
from kalos.core.evaluation import grouped_cv_report
from kalos.core.splits import row_hash_groups
from kalos.core.multiobjective import MultiObjectiveSurrogate, propose_multiobjective
from kalos.core.optimize import propose
from kalos.core.surrogate import DEVICE, DTYPE, FitError, Surrogate
from kalos.data.anonymizer import _hash
from kalos.portal.validate import column_provenance, provenance_dicts

log = logging.getLogger("kalos.portal")

# --- CPU thread budget ------------------------------------------------------- #
# The upload handler offloads its CPU-bound body (pandas parse + ~7 GP fits +
# optimize_acqf) to a worker thread so concurrent uploads do not serialize on the
# event loop. Cap the torch intra-op thread count so several concurrent fits do
# not oversubscribe the CPU and thrash. Override with KALOS_TORCH_THREADS.
_TORCH_THREADS = int(os.environ.get("KALOS_TORCH_THREADS", str(min(4, os.cpu_count() or 1))))
torch.set_num_threads(max(1, _TORCH_THREADS))

app = FastAPI(title="Kalos Engine API")

# --- upload safety limits ---------------------------------------------------- #
# This engine ingests untrusted run sheets from external clients, so the raw
# upload and its parsed shape are capped to bound memory and blunt zip-bomb
# expansion. Sizes are configurable via env; the CSV/xlsx caps are constants.
_MAX_UPLOAD_MB = float(os.environ.get("KALOS_MAX_UPLOAD_MB", "25"))
MAX_UPLOAD_BYTES = int(_MAX_UPLOAD_MB * 1024 * 1024)
MAX_CSV_ROWS = 100_000       # rows read from a CSV/TSV upload
MAX_COLUMNS = 512            # columns allowed in any upload (CSV or xlsx)
MAX_XLSX_CELLS = 2_000_000   # rows * cols ceiling for a parsed xlsx (zip-bomb guard)
_ZIP_MAGIC = b"PK\x03\x04"   # xlsx/xls-as-zip start-of-file marker

# Separate, tighter cap on the rows the SURROGATE is actually fit on. A
# SingleTaskGP is O(n^2) in memory and O(n^3) in time, so a raw upload that is
# within MAX_CSV_ROWS can still be far too large for an exact GP. The optimizer
# targets the small-sample bioprocess regime (N <= a couple thousand), so we
# reject an over-cap fit set rather than silently subsampling (which would be
# invisible, non-deterministic data loss). Override with KALOS_MAX_FIT_ROWS.
MAX_FIT_ROWS = int(os.environ.get("KALOS_MAX_FIT_ROWS", "2000"))

# Generic, non-leaking messages. We never echo the parser error, a column name,
# or a cell value back to an unauthenticated caller.
_ERR_PARSE = "Could not parse the uploaded file. Check it is a CSV or Excel run-sheet."
_ERR_TOO_LARGE = "The uploaded file is too large."
_ERR_TOO_MANY_COLUMNS = "The uploaded file has too many columns."
_ERR_TOO_MANY_FIT_ROWS = (
    "The dataset is too large for the surrogate; the optimizer targets the "
    f"small-sample regime, N<={MAX_FIT_ROWS}."
)
# A numerically-hard-but-valid file (near-duplicate or ill-conditioned rows) that
# defeats even the jittered fit retry gets its OWN message, so it is not confused
# with the generic parse failure.
_ERR_FIT = (
    "The model could not be fit on this data - likely near-duplicate or "
    "ill-conditioned rows."
)


class UploadRejected(ValueError):
    """A client upload failed a safety guard. Carries a generic, safe message."""

# Allow the kalos-web Next.js app (dev + any localhost) to call the engine.
app.add_middleware(
    CORSMiddleware,
    allow_origin_regex=r"http://(localhost|127\.0\.0\.1)(:\d+)?",
    allow_methods=["*"],
    allow_headers=["*"],
)

_HTML = (Path(__file__).parent / "index.html").read_text()
BOUNDS = np.array([[0, 0, 0], [1, 1, 1]], float)
TITER_OPT = np.array([0.7, 0.3, 0.5])
PURITY_OPT = np.array([0.2, 0.8, 0.4])  # different recipe -> titer/purity trade off


# Synthetic but honest demo objectives in real units: titer 0..22 mg/L (peaks at
# TITER_OPT), purity 0..100 % (peaks at a different recipe, so they trade off).
def _titer(X: np.ndarray) -> np.ndarray:
    d = np.sum((np.atleast_2d(X) - TITER_OPT) ** 2, axis=1)
    return 22.0 * np.exp(-3.0 * d)


def _purity(X: np.ndarray) -> np.ndarray:
    d = np.sum((np.atleast_2d(X) - PURITY_OPT) ** 2, axis=1)
    return 100.0 * np.exp(-3.0 * d)


def _objectives(X: np.ndarray) -> np.ndarray:
    x = np.atleast_2d(X)
    return np.stack([_titer(x), _purity(x)], axis=-1)


def _annotate(batch: np.ndarray, mean, std, best: float, cols=None) -> list:
    """Attach predicted value, uncertainty, and an explore/exploit rationale to
    each proposed experiment. Explore = high model uncertainty (chosen to learn);
    exploit = high predicted value (chosen to win). Current human-in-the-loop BO
    research says a recommendation must carry exactly this."""
    mean = np.asarray(mean, float).reshape(-1)
    std = np.asarray(std, float).reshape(-1)
    thr = float(np.quantile(std, 2 / 3)) if len(std) > 2 else float(std.max() if len(std) else 0.0)
    eps = 0.02 * max(abs(best), 1e-9)
    rows = []
    for i in range(len(batch)):
        m, sd = float(mean[i]), float(std[i])
        gain = m - best
        if gain > eps:
            mode, reason = "exploit", f"predicted high (+{gain:.3g} vs best)"
        elif sd >= thr:
            mode, reason = "explore", f"reduce model uncertainty here (±{sd:.3g})"
        else:
            mode, reason = "explore", f"diversifies the batch (predicted {m:.3g}, ±{sd:.3g})"
        row = {"pred": round(m, 3), "std": round(sd, 3), "mode": mode, "reason": reason}
        row["vals"] = (np.round(batch[i], 3).tolist() if cols is None
                       else [round(float(batch[i][j]), 3) for j in cols])
        rows.append(row)
    return rows


def _compute_embedding(X: np.ndarray, y: np.ndarray, seed: int) -> dict | None:
    """A 2D embedding of the observed runs (client-facing scatter, colored by
    target), so a dashboard can show where uploaded runs sit relative to each
    other in feature space, not just the driver table.

    Features are z-scored first (`StandardScaler`), then UMAP is tried; if
    umap-learn is not installed or fitting it fails for any reason, this falls
    back to PCA (always available via scikit-learn). `method` records whichever
    ACTUALLY ran, never hardcoded, so a "umap" label is never shown for a PCA
    projection. Seeded with the same `seed` as the rest of `_analyze` for
    reproducibility. Only computed with enough runs/features to be a meaningful
    2D layout (n >= 5, d >= 2); returns None otherwise so the caller omits the
    key rather than shipping a degenerate embedding."""
    n, d = X.shape
    if n < 5 or d < 2:
        return None
    Xs = StandardScaler().fit_transform(X)
    coords = None
    method = "pca"
    try:
        import umap

        reducer = umap.UMAP(n_components=2, n_neighbors=min(15, n - 1), random_state=seed)
        coords = reducer.fit_transform(Xs)
        method = "umap"
    except Exception:  # noqa: BLE001 - any umap unavailable/runtime failure falls back to PCA
        coords = None
    if coords is None:
        coords = PCA(n_components=2, random_state=seed).fit_transform(Xs)
        method = "pca"
    points = [
        [round(float(coords[i, 0]), 4), round(float(coords[i, 1]), 4), round(float(y[i]), 4)]
        for i in range(n)
    ]
    return {"method": method, "points": points}


def _feature_correlation(feat_frame: pd.DataFrame, y: np.ndarray, feats: list, target: str) -> dict | None:
    """Spearman correlation among the kept features and the target - the same rho
    measure the drivers use, so the heatmap is consistent with the driver table.
    Symmetric matrix, features first then the target. None if too few features."""
    if len(feats) < 2:
        return None
    frame = feat_frame.copy()
    frame[target] = y
    labels = [*feats, target]
    corr = frame.corr(method="spearman").reindex(index=labels, columns=labels)
    matrix = [[0.0 if v != v else round(float(v), 3) for v in row] for row in corr.to_numpy()]
    return {"labels": [str(c) for c in labels], "matrix": matrix}


def _response_surface(s, X: np.ndarray, y: np.ndarray, feats: list, drv: list, grid: int = 24) -> dict | None:
    """GP-predicted target across the two strongest drivers, other inputs held at
    their median. A MODEL prediction (surrogate posterior mean) surfaced so the UI
    can show the predicted landscape - the UI labels it a prediction, not a
    measurement. `runs` are the observed rows projected onto the two axes for
    overlay. None if fewer than two drivers/features or the grid prediction fails."""
    if len(feats) < 2 or len(drv) < 2:
        return None
    try:
        xi, yi = feats.index(drv[0][0]), feats.index(drv[1][0])
        x_vals = np.linspace(X[:, xi].min(), X[:, xi].max(), grid)
        y_vals = np.linspace(X[:, yi].min(), X[:, yi].max(), grid)
        gx, gy = np.meshgrid(x_vals, y_vals)
        pts = np.tile(np.median(X, axis=0), (grid * grid, 1))
        pts[:, xi], pts[:, yi] = gx.ravel(), gy.ravel()
        z = np.asarray(s.posterior(pts)[0], float).reshape(grid, grid)
    except Exception:  # noqa: BLE001 - a degenerate surface just omits the optional field
        return None
    return {
        "x_feature": str(drv[0][0]),
        "y_feature": str(drv[1][0]),
        "x_vals": [round(float(v), 4) for v in x_vals],
        "y_vals": [round(float(v), 4) for v in y_vals],
        "z": [[round(float(v), 4) for v in row] for row in z],
        "runs": [
            [round(float(X[i, xi]), 4), round(float(X[i, yi]), 4), round(float(y[i]), 4)]
            for i in range(X.shape[0])
        ],
    }


# --- persistence of the most-recently analyzed real dataset ------------------ #
# The Overview reads /api/latest so the landing page reflects the LAST dataset a
# user actually uploaded, not the synthetic demo objective. In-memory is the
# source of truth; the JSON file is best-effort so it survives a portal restart.
_STATE_DIR = Path(os.environ.get("KALOS_STATE_DIR", Path.home() / ".kalos"))
_LATEST_PATH = _STATE_DIR / "latest_analysis.json"
_LATEST: dict | None = None
# `_analyze` now runs in a worker thread (see `run_uploaded`), so `_save_latest`
# and `_load_latest` touch the `_LATEST` global and the state file from multiple
# threads concurrently. This lock serializes the read-modify-write + file write
# so two overlapping uploads cannot interleave a half-written global or file.
_LATEST_LOCK = threading.Lock()


def _load_latest() -> dict | None:
    global _LATEST
    with _LATEST_LOCK:
        if _LATEST is None and _LATEST_PATH.exists():
            try:
                _LATEST = json.loads(_LATEST_PATH.read_text())
            except (OSError, ValueError):  # a corrupt cache must not break the API
                _LATEST = None
        return _LATEST


def _save_latest(result: dict, dataset: str) -> None:
    global _LATEST
    with _LATEST_LOCK:
        _LATEST = {**result, "dataset": dataset, "updated": time.time()}
        try:
            _STATE_DIR.mkdir(parents=True, exist_ok=True)
            _LATEST_PATH.write_text(json.dumps(_LATEST))
        except (OSError, TypeError, ValueError):
            # Disk persistence is best-effort; in-memory still serves this run and a
            # serialization hiccup must never fail the upload.
            pass


@app.get("/api/latest")
def latest() -> dict:
    """The most recent real (uploaded) analysis, for the Overview. `has_data` is
    False until the first successful /api/run so the home can show an upload
    prompt instead of pretending there is data."""
    data = _load_latest()
    if not data:
        return {"has_data": False}
    return {"has_data": True, **data}


@app.get("/", response_class=HTMLResponse)
def index() -> str:
    return _HTML


@app.get("/api/single")
def run_single(rounds: int = 6, q: int = 2) -> dict:
    rng = np.random.default_rng(0)
    X = rng.uniform(0, 1, (6, 3))
    y = _titer(X)
    traj = [{"round": 0, "best": float(y.max()), "n": int(len(y))}]
    proposals: list = []
    for r in range(rounds):
        s = Surrogate().fit(X, y, bounds=BOUNDS)
        nxt = propose(s, BOUNDS, q=q)
        X = np.vstack([X, nxt])
        y = np.concatenate([y, _titer(nxt)])
        traj.append({"round": r + 1, "best": float(y.max()), "n": int(len(y))})
    # annotate the last proposed batch with predicted value + uncertainty + why
    incumbent = float(y[:-q].max())  # best before this batch was measured
    mean, std = s.posterior(nxt)
    proposals = _annotate(nxt, mean, std, incumbent)
    bi = int(y.argmax())
    return {
        "trajectory": traj,
        "proposals": proposals,
        "best_recipe": np.round(X[bi], 3).tolist(),
        "best_value": float(y[bi]),
        "optimum": TITER_OPT.tolist(),
    }


@app.get("/api/multi")
def run_multi(rounds: int = 5, q: int = 2) -> dict:
    rng = np.random.default_rng(0)
    X = rng.uniform(0, 1, (8, 3))
    Y = _objectives(X)
    traj: list = []
    proposals: list = []
    for r in range(rounds):
        s = MultiObjectiveSurrogate().fit(X, Y, bounds=BOUNDS)
        nxt = propose_multiobjective(s, BOUNDS, q=q)
        X = np.vstack([X, nxt])
        Y = np.vstack([Y, _objectives(nxt)])
        _, py = s.pareto()
        traj.append(
            {"round": r + 1, "pareto": int(len(py)), "best_titer": float(Y[:, 0].max()), "best_purity": float(Y[:, 1].max())}
        )
        last_batch = nxt
    s = MultiObjectiveSurrogate().fit(X, Y, bounds=BOUNDS)
    # predicted titer + purity (with uncertainty) for each proposed experiment
    import torch
    post = s.model.posterior(torch.as_tensor(last_batch, dtype=DTYPE, device=DEVICE))
    mY = post.mean.detach().cpu().numpy().reshape(len(last_batch), -1)
    sdY = post.variance.clamp_min(1e-12).sqrt().detach().cpu().numpy().reshape(len(last_batch), -1)
    sd_mean = sdY.mean(axis=1)
    sd_thr = float(np.quantile(sd_mean, 2 / 3)) if len(sd_mean) > 2 else float(sd_mean.max())
    proposals = [{
        "vals": np.round(last_batch[i], 3).tolist(),
        "pred_titer": round(float(mY[i, 0]), 1), "pred_purity": round(float(mY[i, 1]), 0),
        "mode": "explore" if sd_mean[i] >= sd_thr else "exploit",
        "reason": "fills a gap on the titer/purity frontier",
    } for i in range(len(last_batch))]
    _, py = s.pareto()
    order = np.argsort(py[:, 0])
    return {
        "trajectory": traj,
        "pareto": [{"titer": float(py[i, 0]), "purity": float(py[i, 1])} for i in order],
        "all": [{"titer": float(Y[i, 0]), "purity": float(Y[i, 1])} for i in range(len(Y))],
        "proposals": proposals,
    }


# --- run the engine on an uploaded data file --------------------------------- #
# Outcome-like columns are candidate TARGETS and are excluded from input features
# so the model never "predicts" titer from another measured output (leakage).
_OUTCOME_HINT = re.compile(r"titer|titre|yield|conc|purity|lipase|biomass|od\d|product|response|output|score|kda|activity|titer", re.I)
_TARGET_PREF = re.compile(r"titer|titre|lipase|yield", re.I)
_ID_HINT = re.compile(r"^(id|name|sample.*|well|index|run|experiment|round|date|time|medium|strain|recipe|batch|campaign|group|lot|notes?)$", re.I)
_GROUP_HINT = re.compile(r"medium|strain|recipe|batch|campaign|group|lot", re.I)


def _numeric_cols(df: pd.DataFrame) -> list:
    # Numeric if >=80% of NON-BLANK cells parse as numbers. Blanks are treated as
    # "absent" (filled with 0 later), so a sparse component column still counts.
    out = []
    for c in df.columns:
        s = df[c]
        nonblank = s.notna() & (s.astype(str).str.strip() != "")
        if nonblank.sum() < 3:
            continue
        if pd.to_numeric(s[nonblank], errors="coerce").notna().mean() >= 0.8:
            out.append(c)
    return out


# Deterministic seed for the analyze path. The same upload -> the same GP fit and
# the same proposed batch, which matters for client reproducibility and audit.
ANALYZE_SEED = 1234


def _seed_everything(seed: int = ANALYZE_SEED) -> None:
    """Seed torch + numpy so one upload yields one deterministic set of proposals."""
    torch.manual_seed(seed)
    np.random.seed(seed)


def _dedupe_columns(df: pd.DataFrame) -> pd.DataFrame:
    """Suffix duplicate column labels (`X`, `X.1`, ...) so every column is a Series.

    A run sheet with a repeated header would otherwise make `df[label]` return a
    2-D frame and break the analysis. pandas already does this on CSV read; we do
    it here too so a directly-built frame (or an xlsx with duplicate headers) is
    handled identically, and the duplicate stays visible in the provenance report.
    """
    if not df.columns.duplicated().any():
        return df
    seen: dict[str, int] = {}
    new_cols: list[str] = []
    for col in df.columns:
        key = str(col)
        if key in seen:
            seen[key] += 1
            new_cols.append(f"{key}.{seen[key]}")
        else:
            seen[key] = 0
            new_cols.append(key)
    out = df.copy()
    out.columns = new_cols
    return out


def _analyze(
    df: pd.DataFrame, target: str | None = None, *, anonymize: bool = False
) -> dict:
    """Run the engine on an arbitrary run sheet: pick the target (the value to
    maximize), use the process INPUTS as features (other measured outputs are
    excluded to avoid leakage), then honest grouped-CV, signed drivers, and a
    proposed next batch.

    Deterministic: seeds torch + numpy up front so the same sheet gives the same
    proposals. Returns a per-column `provenance` report (what was kept/dropped and
    why) plus `seed`, `timestamp`, and `engine_version` for audit. When
    `anonymize` is True, identifier-type column names are replaced with stable
    pseudonyms in the response (the owner UI keeps real names when False)."""
    _seed_everything()
    df = _dedupe_columns(df.dropna(axis=1, how="all"))
    num = _numeric_cols(df)
    if not num:
        raise ValueError("no numeric columns found")
    outcomes = [c for c in num if _OUTCOME_HINT.search(str(c))]
    if target is None or target not in df.columns:
        target = next((c for c in outcomes if _TARGET_PREF.search(str(c))), outcomes[0] if outcomes else num[-1])
    feats = []
    for c in num:
        if c == target or _ID_HINT.match(str(c).strip()) or _OUTCOME_HINT.search(str(c)):
            continue  # drop the target, ids, and OTHER measured outputs (anti-leakage)
        col = pd.to_numeric(df[c], errors="coerce")
        if col.std(skipna=True) and col.std() > 1e-9:
            feats.append(c)
    if len(feats) < 1:
        raise ValueError("no varying process-input columns found (only outputs/ids?)")
    candidate_targets = outcomes or [target]

    y_all = pd.to_numeric(df[target], errors="coerce")
    keep = y_all.notna()
    X_raw = df.loc[keep, feats].apply(pd.to_numeric, errors="coerce")   # NaN preserved (for grouping)
    X_zf = X_raw.fillna(0.0)                                             # zero-filled (for the GP + box)
    y = y_all[keep].to_numpy(float)
    if len(y) < 6:
        raise ValueError(f"need at least 6 rows with a numeric {target!r}; got {len(y)}")
    # Cap the rows the O(n^2) exact GP is fit on, separately from the raw-upload
    # cap. Reject rather than subsample: silent subsampling would be invisible,
    # non-deterministic data loss and contradict the reproducibility contract.
    if len(y) > MAX_FIT_ROWS:
        raise UploadRejected(_ERR_TOO_MANY_FIT_ROWS)

    # The design box is built from the target-present (fitted) rows, so the honest
    # varying-feature check must be recomputed on THOSE rows, not the full column. A
    # feature that varies over the whole sheet but is constant on the fitted rows
    # would otherwise collapse its bound to zero width silently. Drop such features
    # and record them so provenance flags them instead of misleading the client.
    fitted_range = X_zf.max(axis=0) - X_zf.min(axis=0)
    constant_on_fitted = [c for c in feats if float(fitted_range[c]) <= 1e-9]
    if constant_on_fitted:
        feats = [c for c in feats if c not in set(constant_on_fitted)]
        if len(feats) < 1:
            raise ValueError("no varying process-input columns on the target-present rows")
        X_raw = X_raw.drop(columns=constant_on_fitted)
        X_zf = X_zf.drop(columns=constant_on_fitted)

    X = X_zf.to_numpy(float)
    bounds = np.vstack([X.min(0), X.max(0)])
    embedding = _compute_embedding(X, y, ANALYZE_SEED)

    gcol = next((c for c in df.columns if _GROUP_HINT.search(str(c))), None)
    if gcol:
        groups = df.loc[keep, gcol].astype(str).tolist()
    else:
        # group on the RAW values (NaN preserved) so rows missing different
        # components are not merged into one replicate group by the zero-fill —
        # via the one leakage-checked, deterministic grouper.
        groups = row_hash_groups(X_raw)

    # honest grouped cross-validation: pooled out-of-fold predictions + a
    # group-level bootstrap CI, all through the single leakage-checked splitter.
    rep = grouped_cv_report(X, y, groups=groups, n_splits=5, bounds=bounds)
    rho = rep["spearman"]
    oof_a, oof_p = rep["oof_actual"], rep["oof_pred"]

    # distribution-free +/- band from the pooled out-of-fold residuals (approximate
    # coverage under grouped CV). Honest alternative to the surrogate's own std,
    # which is often overconfident on small bioprocess datasets.
    resid = np.asarray(oof_a, float) - np.asarray(oof_p, float)
    conformal_q = round(q_from_residuals(resid, alpha=0.1), 4) if len(resid) else None

    # Honest reliability verdict: only what this path can actually assess. The
    # spearman floor mirrors GatesConfig.min_spearman (kalos/core/gates.py); we do
    # NOT assert feasibility or calibration gates, which are not measured here.
    ci95 = None if rho != rho else [round(rep["ci95"][0], 3), round(rep["ci95"][1], 3)]
    reliability = {
        "spearman": None if rho != rho else round(rho, 3),
        "ci95": ci95,
        "spearman_floor": 0.20,
        "clears_floor": bool(rho == rho and rho >= 0.20),
        "ci_excludes_zero": bool(ci95 is not None and ci95[0] > 0),
        "unmodeled": ["feasibility probability", "calibration (ECE)", "scale-up transfer"],
    }

    # signed drivers
    drv = []
    for c in feats:
        r = spearmanr(pd.to_numeric(df.loc[keep, c], errors="coerce").fillna(0.0), y).statistic
        drv.append([str(c), 0.0 if r != r else float(r)])
    drv.sort(key=lambda d: -abs(d[1]))
    drv = drv[:8]

    # proposed next batch, with predicted target + uncertainty + a why per row
    s = Surrogate().fit(X, y, bounds=bounds)
    batch = propose(s, bounds, q=5)
    show = [d[0] for d in drv[:4]]
    show_idx = [feats.index(f) for f in show]
    p_mean, p_std = s.posterior(batch)
    correlation = _feature_correlation(X_zf, y, feats, str(target))
    surface = _response_surface(s, X, y, feats, drv)

    # per-column provenance: what was kept as a feature, used as the target, or
    # dropped (id / other output / constant / sparse), so the client is never left
    # guessing about a silently dropped column. Mirrors the selection logic above.
    provenance = provenance_dicts(
        column_provenance(
            df,
            target=str(target),
            features=[str(c) for c in feats],
            numeric_cols=[str(c) for c in num],
            id_hint=_ID_HINT,
            outcome_hint=_OUTCOME_HINT,
            constant_on_fitted_rows=[str(c) for c in constant_on_fitted],
        )
    )

    result = {
        "n": int(keep.sum()), "d": len(feats), "target": str(target), "group_col": gcol,
        "targets": [str(c) for c in candidate_targets], "features": [str(c) for c in feats],
        "cv_spearman": None if rho != rho else round(rho, 3),
        "cv_ci95": ci95,
        "cv_n_groups": rep["n_groups"],
        "conformal_q": conformal_q,
        "reliability": reliability,
        "best": round(float(y.max()), 4),
        "drivers": [{"name": c, "rho": round(r, 3)} for c, r in drv],
        "proposal_features": show,
        "proposals": _annotate(batch, p_mean, p_std, float(y.max()), cols=show_idx),
        "oof": [[round(a, 4), round(p, 4)] for a, p in zip(oof_a, oof_p)],
        "provenance": provenance,
        "seed": ANALYZE_SEED,
        "timestamp": int(time.time()),
        "engine_version": ENGINE_VERSION,
    }
    if embedding is not None:
        result["embedding"] = embedding
    if correlation is not None:
        result["correlation"] = correlation
    if surface is not None:
        result["response_surface"] = surface
    if anonymize:
        result = _anonymize_result(result)
    return result


def _anonymize_result(result: dict) -> dict:
    """Replace identifier-type column names in the response with stable pseudonyms.

    Only identifier-like columns (matched by `_ID_HINT`) are renamed; process
    features and the target keep their real names because the authenticated owner
    UI legitimately shows them (a driver like "Methanol"). The mapping is stable
    (same name -> same pseudonym) via the anonymizer's irreversible hash, so an
    anonymized report is still internally consistent across fields.
    """
    def alias(name: str) -> str:
        return f"col_{_hash(name)[:8]}" if _ID_HINT.match(str(name).strip()) else name

    out = dict(result)
    if out.get("group_col"):
        out["group_col"] = alias(out["group_col"])
    out["provenance"] = [
        {**row, "name": alias(row["name"])} for row in out.get("provenance", [])
    ]
    return out


def _reject_oversized_xlsx(raw: bytes) -> None:
    """Reject an xlsx whose declared dimensions exceed the cell / column caps,
    BEFORE `pd.read_excel` materializes the frame (the zip-bomb guard).

    Opens the workbook read-only with openpyxl and reads each sheet's declared
    dimension (`max_row * max_column`) without loading cell values. If any sheet's
    declared cell count exceeds `MAX_XLSX_CELLS` (or its column count exceeds
    `MAX_COLUMNS`), raise `UploadRejected` so the caller never allocates the full
    frame. A corrupt or unreadable zip raises `UploadRejected` too (client error).
    """
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True)
    except Exception as exc:  # noqa: BLE001 - corrupt/non-xlsx zip -> generic 400
        log.warning("rejected an unreadable xlsx upload: %s", type(exc).__name__)
        raise UploadRejected(_ERR_PARSE) from exc
    try:
        for ws in wb.worksheets:
            cols = ws.max_column or 0
            rows = ws.max_row or 0
            if cols > MAX_COLUMNS:
                raise UploadRejected(_ERR_TOO_MANY_COLUMNS)
            if rows * cols > MAX_XLSX_CELLS:
                raise UploadRejected(_ERR_TOO_LARGE)
    finally:
        wb.close()


def _parse_upload(raw: bytes) -> pd.DataFrame:
    """Turn raw upload bytes into a bounded dataframe, or raise `UploadRejected`.

    Guards, in order:
      1. byte-size cap (default 25 MB, env `KALOS_MAX_UPLOAD_MB`);
      2. filetype sniff by MAGIC BYTES, not extension: a zip header (`PK\\x03\\x04`)
         is treated as xlsx/xls, anything else as UTF-8 text/CSV;
      3. shape caps: CSV/TSV rows and a column ceiling, and an xlsx cell-count
         (rows * cols) ceiling to blunt zip-bomb expansion.
    All rejection messages are generic (no parser text, column, or cell echoed).
    """
    if len(raw) > MAX_UPLOAD_BYTES:
        raise UploadRejected(_ERR_TOO_LARGE)

    if raw[:4] == _ZIP_MAGIC:
        # xlsx/xls: enforce the cell-count ceiling BEFORE pd.read_excel fully
        # materializes the frame, so a zip-bomb whose DECLARED sheet dimensions are
        # enormous is rejected without the memory spike of building the DataFrame.
        # A zip header on a corrupt or non-xlsx zip (BadZipFile, openpyxl's
        # InvalidFileException, ValueError) is a client error, not a server one.
        _reject_oversized_xlsx(raw)
        try:
            df = pd.read_excel(io.BytesIO(raw))
        except Exception as exc:  # noqa: BLE001 - normalized to a generic 400 below
            log.warning("rejected an unreadable xlsx upload: %s", type(exc).__name__)
            raise UploadRejected(_ERR_PARSE) from exc
        rows, cols = df.shape
        # Belt-and-suspenders: re-check the materialized shape. The pre-read guard
        # uses the sheet's declared dimensions; this catches a mismatch and the
        # column ceiling on the actual parsed frame.
        if cols > MAX_COLUMNS:
            raise UploadRejected(_ERR_TOO_MANY_COLUMNS)
        if rows * cols > MAX_XLSX_CELLS:
            raise UploadRejected(_ERR_TOO_LARGE)
        return df

    # Otherwise treat as text/CSV. A binary blob that is neither a zip nor valid
    # tabular text will not yield usable numeric columns and is rejected downstream
    # with the same generic parse message.
    text = raw.decode("utf-8", errors="replace")
    if text.strip() == "":
        raise UploadRejected(_ERR_PARSE)
    head = text[:4000]
    sep = "\t" if head.count("\t") > head.count(",") else ","
    # cap columns first (cheap, from the header) before reading the full body
    ncols = pd.read_csv(io.StringIO(text), sep=sep, nrows=0).shape[1]
    if ncols > MAX_COLUMNS:
        raise UploadRejected(_ERR_TOO_MANY_COLUMNS)
    df = pd.read_csv(io.StringIO(text), sep=sep)
    # Fail closed on the row cap rather than silently truncating to the first
    # MAX_CSV_ROWS rows: silent data loss would contradict the safe-errors
    # contract, so an over-cap CSV is rejected like the xlsx cell-cap guard.
    if len(df) > MAX_CSV_ROWS:
        raise UploadRejected(_ERR_TOO_LARGE)
    return df


def _run_uploaded_sync(raw: bytes, target: str | None, anonymize: bool, filename: str) -> dict:
    """The CPU-bound body of an upload: parse -> analyze -> persist -> result.

    This is the heavy, blocking work (pandas parse, ~7 GP fits, optimize_acqf) and
    runs in a worker thread (see `run_uploaded`), NOT on the asyncio event loop, so
    concurrent uploads do not serialize and GET /api/latest never hangs behind a
    fit. It raises `UploadRejected` / `ValueError` / `FitError` / parser errors;
    the async wrapper maps each to the right 400 envelope. Kept fully synchronous
    so it is trivially unit-testable in isolation.
    """
    df = _parse_upload(raw)
    result = _analyze(df, target, anonymize=anonymize)
    _save_latest(result, filename)
    return result


@app.post("/api/run")
async def run_uploaded(
    file: UploadFile = File(...),
    target: str = Form(default=""),
    anonymize: bool = Form(default=False),
) -> JSONResponse:
    raw = await file.read()
    filename = file.filename or "uploaded dataset"
    try:
        # Offload the CPU-bound parse + fit + save to a worker thread so this
        # single-worker service does not block the event loop (and every other
        # request, including GET /api/latest) while a GP fit runs.
        result = await run_in_threadpool(
            _run_uploaded_sync, raw, target or None, anonymize, filename
        )
        return JSONResponse(result)
    except UploadRejected as rej:
        # A guard tripped: the message is already generic and safe to return.
        log.warning("upload rejected: %s", rej)
        return JSONResponse({"error": str(rej)}, status_code=400)
    except FitError:
        # A numerically-hard-but-valid file that defeated the jittered fit retry.
        # Give it a DISTINCT message so it does not masquerade as a parse failure;
        # the exception text itself is already generic, but we do not echo it.
        log.warning("upload rejected: model could not be fit (ill-conditioned rows)")
        return JSONResponse({"error": _ERR_FIT}, status_code=400)
    except (pd.errors.ParserError, pd.errors.EmptyDataError, ValueError, UnicodeError):
        # Log the full traceback server-side for debugging; return a generic
        # message so no parser detail, column name, cell value, or stack trace
        # ever reaches the (unauthenticated) client.
        log.exception("failed to parse or analyze an uploaded file")
        return JSONResponse({"error": _ERR_PARSE}, status_code=400)
    except Exception:  # noqa: BLE001 - normalized to a generic 400 below
        # Catch-all so a fit-time failure that is NOT a plain ValueError still
        # returns the {error} JSON envelope the frontend parses, never FastAPI's
        # default text/plain HTTP 500. This covers torch.linalg.LinAlgError from a
        # GP fit and an AssertionError from the leakage guard, among others. As
        # above, the full traceback is logged server-side and only the generic
        # message reaches the (unauthenticated) client.
        log.exception("failed to parse or analyze an uploaded file")
        return JSONResponse({"error": _ERR_PARSE}, status_code=400)
